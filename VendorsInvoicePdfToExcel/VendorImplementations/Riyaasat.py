from openpyxl.styles.builtins import normal
import re

from VendorsInvoicePdfToExcel.helper import indexOfContainsInList, convert_to_ddmmyy
from VendorsInvoicePdfToExcel.helper import convert_amount_to_words
from fastapi import HTTPException

class Riyaasat:
    def __init__(self, tables, text_data, table_by_tabula):
        self.tables = tables
        self.text_data = text_data
        self.table_by_tabula = table_by_tabula

    def getVendorInfo(self):
        firstPageText = self.tables[1]
        vendorInfo = firstPageText[indexOfContainsInList(firstPageText, "RIYAA")][0].split("\n")
        return {
            "vendor_name": vendorInfo[0],
            "vendor_address": ", ".join(vendorInfo[:4]),
            "vendor_mob": "N/A",
            "vendor_gst": vendorInfo[indexOfContainsInList(vendorInfo, "GST")].split(":")[-1],
            "vendor_email": "N/A"
        }

    def getInvoiceInfo(self):
        firstPageText = self.tables[1]
        print(convert_to_ddmmyy(self.text_data[1].split("\n")[indexOfContainsInList(self.text_data[1].split("\n"), "Dated") + 1].split(" ")[             -1].strip()))
        return {
            "invoice_number": firstPageText[indexOfContainsInList(firstPageText, "Invoice")][
                indexOfContainsInList(firstPageText[indexOfContainsInList(firstPageText, "Invoice")], "Invoice")].split(
                "\n")[-1],
            "invoice_date": convert_to_ddmmyy(
                self.text_data[1].split("\n")[indexOfContainsInList(self.text_data[1].split("\n"), "Dated") + 1].split(
                    " ")[-1].strip())
        }

    def getReceiverInfo(self):
        firstPageText = self.tables[1]

        if indexOfContainsInList(firstPageText, "Ship to") == -1:
            return {
                "receiver_name": "N/A",
                "receiver_address": "N/A",
                "receiver_gst": "N/A"
            }

        receiverInfo = firstPageText[indexOfContainsInList(firstPageText, "Ship to")][0].split("\n")
        receiverInfo = receiverInfo[indexOfContainsInList(receiverInfo, "Ship") : indexOfContainsInList(receiverInfo, "Bill")]
        return {
            "receiver_name": receiverInfo[1],
            "receiver_address": receiverInfo[2],
            "receiver_gst": receiverInfo[indexOfContainsInList(receiverInfo, "GST")].split(":")[-1].strip()
        }

    def getBillingInfo(self):
        firstPageText = self.tables[1]
        billToInfo = firstPageText[indexOfContainsInList(firstPageText, "Bill to")][0].split("\n")
        billToInfo = billToInfo[indexOfContainsInList(billToInfo, "Bill") :]
        return {
            "billto_name": billToInfo[1],
            "billto_address": billToInfo[2],
            "place_of_supply": billToInfo[4].split(":")[-2].split(",")[0],
            "billto_gst": billToInfo[indexOfContainsInList(billToInfo, "GST")].split(":")[-1].strip()
        }

    def getItemInfo(self):
        pages = self.tables
        firstPage = self.tables[1]
        lastPage = pages[len(pages)]
        total_tax = { "IGST": lastPage[indexOfContainsInList(lastPage, "Taxable")+2][1],  "SGST": 0, "CGST": 0,}

        gstType = ""
        if indexOfContainsInList(lastPage[indexOfContainsInList(lastPage, "Amount Charg") + 1], "CGST") != -1:
            raise HTTPException(status_code=400, detail="For Riyaasat CGST is not implemented")

        if indexOfContainsInList(lastPage[indexOfContainsInList(lastPage, "Amount Charg") + 1], "IGST") != -1:
            gstType = "IGST"

        if indexOfContainsInList(lastPage[indexOfContainsInList(lastPage, "Amount Charg") + 1], "SGST") != -1:
            raise HTTPException(status_code=400, detail="For Riyaasat SGST is not implemented")


        products = []
        indexOfHeader = indexOfContainsInList(self.tables[1], "HSN/")
        indexOfSr = indexOfContainsInList(firstPage[indexOfHeader], "Sl")
        indexOfItemname = indexOfContainsInList(firstPage[indexOfHeader], "Description")
        indexOfHsn = indexOfContainsInList(firstPage[indexOfHeader], "HSN")
        indexOfQty = indexOfContainsInList(firstPage[indexOfHeader], "Quantity")
        indexOfPer = indexOfContainsInList(firstPage[indexOfHeader], "per")
        indexOfRate = indexOfContainsInList(firstPage[indexOfHeader], "Rate")
        indexOfAmt = indexOfContainsInList(firstPage[indexOfHeader], "Amount")

        for itemIndex, item in enumerate(firstPage[indexOfHeader+1:]):
            if indexOfContainsInList(item, "Amount") is not -1:
                break
            if item[0].strip() == "":
                continue

            aProductResult= {}
            aProductResult["po_no"] = ""
            aProductResult["or_po_no"] = ""
            orPoInfo = firstPage[indexOfContainsInList(firstPage, "Remark")][0].split("\n")[indexOfContainsInList(firstPage[indexOfContainsInList(firstPage, "Remark")][0].split("\n"), "Remark")+1].split(":")[-1]
            if orPoInfo.find("OR") is not -1:
                aProductResult["or_po_no"] = orPoInfo
            else:
                aProductResult["po_no"] = orPoInfo

            gstPercenatge = item[indexOfItemname][item[indexOfItemname].find("@")+2 : item[indexOfItemname].find("%")]
            aProductResult["debit_note_no"] = ""
            aProductResult["index"] =  item[indexOfSr]
            aProductResult["vendor_code"] = ""
            aProductResult["HSN/SAC"] = item[indexOfHsn]
            aProductResult["Qty"] = item[indexOfQty]
            aProductResult["Rate"] = item[indexOfRate]
            aProductResult["Per"] = item[indexOfPer]
            aProductResult["mrp"] = item[indexOfRate]
            aProductResult["Amount"] = item[indexOfAmt].split("\n")[0]
            aProductResult["po_cost"] = ""
            aProductResult["gst_rate"] = item[indexOfItemname][item[indexOfItemname].find("@")+2 : item[indexOfItemname].find("%")]
            aProductResult["gst_type"] = gstType
            aProductResult["tax_applied"] = float(item[indexOfAmt].split("\n")[0].replace(",","")) * float(gstPercenatge) * 0.01
            products.append(aProductResult)

        return products, total_tax


    def getVendorBankInfo(self):
        return {
            "bank_name": "",
            "account_number": "",
            "ifs_code": "",
        }

    def getItemTotalInfo(self):
        lastPage  =self.tables[len(self.tables)]
        returnData = {}
        returnData["tax_amount_in_words"] = lastPage[indexOfContainsInList(lastPage, "Tax Amount (")][0].split("\n")[0].split(":")[-1]
        returnData["amount_charged_in_words"] = lastPage[indexOfContainsInList(lastPage, "Amount Ch")][0].split("\n")[-1]
        returnData["total_pcs"] = lastPage[indexOfContainsInList(lastPage, "Total")][3]
        returnData["total_amount_after_tax"] = lastPage[indexOfContainsInList(lastPage, "Total")][-1].split(" ")[-1]
        returnData["total_b4_tax"] = lastPage[indexOfContainsInList(lastPage, "Taxable")+2][1]
        returnData["total_tax"] =lastPage[indexOfContainsInList(lastPage, "Taxable")+2][-1]
        returnData["tax_rate"] = lastPage[indexOfContainsInList(lastPage, "Taxable")+2][2].replace("%", "")
        returnData["total_tax_percentage"] =returnData["tax_rate"]
        return returnData